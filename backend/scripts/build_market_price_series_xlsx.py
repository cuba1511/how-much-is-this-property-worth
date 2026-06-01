"""XLSX → SQLite preprocessor for the market price series.

Reads ``backend/data/data.xlsx`` (TF Labs snapshot export, ~8 000 rows with
current + y1ago/y2ago/y3ago columns) and produces the same SQLite schema as
:mod:`build_market_price_series` so the rest of the codebase reads it
transparently.

Period mapping (assuming Q1 2026 snapshot):
    sale_price_sqm       → 2026-01
    sale_price_sqm_y1ago → 2025-01
    sale_price_sqm_y2ago → 2024-01
    sale_price_sqm_y3ago → 2023-01

Re-run with:
    python backend/scripts/build_market_price_series_xlsx.py [--xlsx PATH] [--db PATH]
"""

from __future__ import annotations

import argparse
import logging
import os
import sqlite3
import sys
import time
import unicodedata
from pathlib import Path

logger = logging.getLogger("build_market_price_series_xlsx")

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_XLSX = REPO_ROOT / "backend" / "data" / "data.xlsx"
DEFAULT_DB = REPO_ROOT / "backend" / "data" / "market_price_series.db"

PERIOD_MAP = {
    "sale_price_sqm": "2026-01",
    "sale_price_sqm_y1ago": "2025-01",
    "sale_price_sqm_y2ago": "2024-01",
    "sale_price_sqm_y3ago": "2023-01",
}


def _normalize_name(name: str) -> str:
    decomposed = unicodedata.normalize("NFKD", name)
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch)).lower().strip()


def build(xlsx_path: Path, db_path: Path) -> dict[str, int | str | float]:
    try:
        import openpyxl
    except ImportError:
        logger.error("openpyxl is required: pip install openpyxl")
        raise

    if not xlsx_path.exists():
        raise FileNotFoundError(f"XLSX not found at {xlsx_path}.")

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        raise ValueError("No active sheet in workbook")

    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {name: idx for idx, name in enumerate(headers) if name}

    required = {"town_code", "town_name", "sale_price_sqm"}
    missing = required - set(col)
    if missing:
        raise ValueError(f"Missing columns: {missing}")

    db_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = db_path.with_suffix(db_path.suffix + ".tmp")
    if tmp_path.exists():
        tmp_path.unlink()

    started = time.monotonic()
    conn = sqlite3.connect(tmp_path)

    try:
        conn.execute("PRAGMA journal_mode = OFF")
        conn.execute("PRAGMA synchronous = OFF")
        conn.executescript("""
            DROP TABLE IF EXISTS market_price_series;
            DROP TABLE IF EXISTS market_towns;
            CREATE TABLE market_towns (
                town_id        TEXT PRIMARY KEY,
                ine_code       TEXT NOT NULL,
                town_name      TEXT NOT NULL,
                town_name_norm TEXT NOT NULL,
                province_id    TEXT,
                community_id   TEXT
            );
            CREATE TABLE market_price_series (
                town_id TEXT NOT NULL REFERENCES market_towns(town_id),
                period  TEXT NOT NULL,
                value   REAL NOT NULL,
                PRIMARY KEY (town_id, period)
            ) WITHOUT ROWID;
        """)

        towns_inserted = 0
        series_inserted = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            ine_code = str(row[col["town_code"]]).strip() if row[col["town_code"]] else ""
            town_name = str(row[col["town_name"]]).strip() if row[col["town_name"]] else ""
            if not ine_code or not town_name:
                continue

            town_id = ine_code
            province = str(row[col.get("province_name", -1)]).strip() if col.get("province_name") is not None and row[col["province_name"]] else None
            community = str(row[col.get("autonomus_community_name", -1)]).strip() if col.get("autonomus_community_name") is not None and row[col["autonomus_community_name"]] else None

            conn.execute(
                "INSERT OR IGNORE INTO market_towns "
                "(town_id, ine_code, town_name, town_name_norm, province_id, community_id) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (town_id, ine_code, town_name, _normalize_name(town_name), province, community),
            )
            towns_inserted += 1

            for col_name, period in PERIOD_MAP.items():
                if col_name not in col:
                    continue
                raw = row[col[col_name]]
                if raw is None:
                    continue
                try:
                    value = float(raw)
                except (TypeError, ValueError):
                    continue
                if value <= 0:
                    continue
                conn.execute(
                    "INSERT OR REPLACE INTO market_price_series (town_id, period, value) "
                    "VALUES (?, ?, ?)",
                    (town_id, period, value),
                )
                series_inserted += 1

        conn.executescript("""
            CREATE INDEX IF NOT EXISTS idx_market_ine_code  ON market_towns (ine_code);
            CREATE INDEX IF NOT EXISTS idx_market_town_name ON market_towns (town_name_norm);
        """)
        conn.commit()
        conn.execute("VACUUM")

        towns = conn.execute("SELECT COUNT(*) FROM market_towns").fetchone()[0]
        first_period, last_period = conn.execute(
            "SELECT MIN(period), MAX(period) FROM market_price_series"
        ).fetchone()
    finally:
        conn.close()

    wb.close()

    db_path.unlink(missing_ok=True)
    os.replace(tmp_path, db_path)

    elapsed = time.monotonic() - started
    return {
        "towns": towns,
        "series_rows": series_inserted,
        "first_period": first_period,
        "last_period": last_period,
        "elapsed_seconds": round(elapsed, 2),
        "db_path": str(db_path),
        "db_size_mb": round(db_path.stat().st_size / (1024 * 1024), 1),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    try:
        stats = build(args.xlsx, args.db)
    except FileNotFoundError as exc:
        logger.error("%s", exc)
        return 2

    logger.info(
        "Built %s towns / %s series rows covering %s → %s in %.2fs (%.1f MB) at %s",
        f"{stats['towns']:,}",
        f"{stats['series_rows']:,}",
        stats["first_period"],
        stats["last_period"],
        stats["elapsed_seconds"],
        stats["db_size_mb"],
        stats["db_path"],
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
